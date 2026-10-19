import csv
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import almacen

ENCABEZADOS = ["#", "Fecha", "Metodo", "Nota", "Abono"]
AZUL = "1F4E78"
AZUL_CLARO = "D9E1F2"
VERDE = "548235"
ROJO = "C00000"


def _fecha_excel(iso):
    try:
        return datetime.strptime(iso, "%Y-%m-%d").date()
    except (ValueError, TypeError):
        return iso


def _datos(abonos, titular, deuda_inicial, moneda):
    if abonos is None:
        abonos = almacen.list_abonos()
    if titular is None:
        titular = almacen.get_config("titular", "Tia")
    if deuda_inicial is None:
        deuda_inicial = almacen.get_deuda_inicial()
    if moneda is None:
        moneda = almacen.get_config("moneda", "MXN")
    return abonos, titular, deuda_inicial, moneda


def exportar_excel(ruta, abonos=None, titular=None, deuda_inicial=None, moneda=None):
    abonos, titular, deuda_inicial, moneda = _datos(abonos, titular, deuda_inicial, moneda)
    fmt_moneda = f'"$"#,##0.00 "{moneda}"'

    wb = Workbook()
    ws = wb.active
    ws.title = "Abonos"

    fino = Side(style="thin", color="BFBFBF")
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)

    ws.merge_cells("A1:E1")
    c = ws["A1"]
    c.value = f"Registro de pagos - {titular}"
    c.font = Font(bold=True, size=16, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=AZUL)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:E2")
    c = ws["A2"]
    c.value = f"Generado el {date.today().strftime('%d/%m/%Y')}"
    c.font = Font(italic=True, size=9, color="808080")
    c.alignment = Alignment(horizontal="center")

    fila_enc = 4
    for col, titulo in enumerate(ENCABEZADOS, start=1):
        c = ws.cell(row=fila_enc, column=col, value=titulo)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = borde

    primera = fila_enc + 1
    fila = primera
    for i, abono in enumerate(abonos, start=1):
        ws.cell(row=fila, column=1, value=i).alignment = Alignment(horizontal="center")
        cf = ws.cell(row=fila, column=2, value=_fecha_excel(abono["fecha"]))
        cf.number_format = "DD/MM/YYYY"
        cf.alignment = Alignment(horizontal="center")
        ws.cell(row=fila, column=3, value=abono.get("metodo", ""))
        ws.cell(row=fila, column=4, value=abono.get("nota", ""))
        cm = ws.cell(row=fila, column=5, value=float(abono["monto"]))
        cm.number_format = fmt_moneda
        if i % 2 == 0:
            for col in range(1, 6):
                ws.cell(row=fila, column=col).fill = PatternFill("solid", fgColor=AZUL_CLARO)
        for col in range(1, 6):
            ws.cell(row=fila, column=col).border = borde
        fila += 1

    ultima = fila - 1
    colm = get_column_letter(5)
    rango = f"{colm}{primera}:{colm}{ultima}" if abonos else f"{colm}{primera}:{colm}{primera}"

    f_total = fila + 1
    ws.cell(row=f_total, column=4, value="Total abonado").font = Font(bold=True)
    ws.cell(row=f_total, column=4).alignment = Alignment(horizontal="right")
    ct = ws.cell(row=f_total, column=5, value=f"=SUM({rango})")
    ct.font = Font(bold=True, color=VERDE)
    ct.number_format = fmt_moneda
    ct.fill = PatternFill("solid", fgColor="E2EFDA")

    f_cuenta = f_total + 1
    ws.cell(row=f_cuenta, column=4, value="Numero de abonos").font = Font(bold=True)
    ws.cell(row=f_cuenta, column=4).alignment = Alignment(horizontal="right")
    ws.cell(row=f_cuenta, column=5, value=f"=COUNT({rango})").alignment = Alignment(
        horizontal="center"
    )

    f_prom = f_cuenta + 1
    ws.cell(row=f_prom, column=4, value="Abono promedio").font = Font(bold=True)
    ws.cell(row=f_prom, column=4).alignment = Alignment(horizontal="right")
    cp = ws.cell(row=f_prom, column=5, value=f"=IF(COUNT({rango})=0,0,AVERAGE({rango}))")
    cp.number_format = fmt_moneda

    f_deuda = f_prom + 2
    ws.cell(row=f_deuda, column=4, value="Deuda inicial").font = Font(bold=True)
    ws.cell(row=f_deuda, column=4).alignment = Alignment(horizontal="right")
    cd = ws.cell(row=f_deuda, column=5, value=float(deuda_inicial))
    cd.number_format = fmt_moneda

    f_saldo = f_deuda + 1
    ws.cell(row=f_saldo, column=4, value="Saldo pendiente").font = Font(bold=True)
    ws.cell(row=f_saldo, column=4).alignment = Alignment(horizontal="right")
    cs = ws.cell(row=f_saldo, column=5, value=f"={colm}{f_deuda}-{colm}{f_total}")
    cs.font = Font(bold=True, color=ROJO)
    cs.number_format = fmt_moneda
    cs.fill = PatternFill("solid", fgColor="FCE4D6")

    for col, ancho in {"A": 6, "B": 14, "C": 16, "D": 28, "E": 20}.items():
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = f"A{primera}"

    wb.save(ruta)
    return ruta


def exportar_csv(ruta, abonos=None, titular=None, deuda_inicial=None, moneda=None):
    abonos, titular, deuda_inicial, moneda = _datos(abonos, titular, deuda_inicial, moneda)
    total = round(sum(float(a["monto"]) for a in abonos), 2)
    saldo = round(float(deuda_inicial) - total, 2)

    with open(ruta, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow([f"Registro de pagos - {titular}"])
        w.writerow([f"Moneda: {moneda}"])
        w.writerow([])
        w.writerow(["#", "Fecha", "Metodo", "Nota", "Abono"])
        for i, a in enumerate(abonos, start=1):
            w.writerow(
                [i, a["fecha"], a.get("metodo", ""), a.get("nota", ""), f"{float(a['monto']):.2f}"]
            )
        w.writerow([])
        w.writerow(["", "", "", "Total abonado", f"{total:.2f}"])
        w.writerow(["", "", "", "Numero de abonos", len(abonos)])
        w.writerow(["", "", "", "Deuda inicial", f"{float(deuda_inicial):.2f}"])
        w.writerow(["", "", "", "Saldo pendiente", f"{saldo:.2f}"])
    return ruta
